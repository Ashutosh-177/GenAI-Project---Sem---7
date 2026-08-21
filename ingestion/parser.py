"""Routes a document to the right extractor by file type and returns a flat
list of page/slide/sheet-level text units with source metadata attached.
Nothing here chunks or embeds — that's retrieval/'s job. This module's only
concern is: given a file, get clean text out of it, however it's stored."""
from dataclasses import dataclass
from pathlib import Path

import pdfplumber
import pymupdf
import pytesseract
from PIL import Image
from docx import Document as DocxDocument
from openpyxl import load_workbook
from pptx import Presentation

from config import TESSERACT_CMD, OCR_FALLBACK_CHAR_THRESHOLD

pytesseract.pytesseract.tesseract_cmd = TESSERACT_CMD


@dataclass
class PageUnit:
    source: str          # filename
    unit_label: str       # "page 3", "slide 2", "sheet Sheet1", etc.
    text: str
    extraction_method: str  # "native" | "ocr"


def parse_document(path: str | Path) -> list[PageUnit]:
    path = Path(path)
    ext = path.suffix.lower()
    if ext == ".pdf":
        return _parse_pdf(path)
    elif ext == ".docx":
        return _parse_docx(path)
    elif ext in (".xlsx", ".xlsm"):
        return _parse_xlsx(path)
    elif ext in (".pptx", ".ppt"):
        return _parse_pptx(path)
    else:
        raise ValueError(f"Unsupported file type: {ext} ({path.name})")


def _parse_pdf(path: Path) -> list[PageUnit]:
    """Try native text extraction per page first (fast, accurate). Any page
    that comes back near-empty is assumed to be a scan/image and gets OCR'd
    instead — this is what lets one pipeline handle both digital PDFs and
    scanned ones without the caller having to know which is which."""
    units = []
    with pdfplumber.open(path) as pdf:
        doc = pymupdf.open(path)  # opened alongside for OCR fallback rendering
        for i, page in enumerate(pdf.pages):
            native_text = (page.extract_text() or "").strip()
            if len(native_text) >= OCR_FALLBACK_CHAR_THRESHOLD:
                units.append(PageUnit(path.name, f"page {i + 1}", native_text, "native"))
            else:
                ocr_text = _ocr_pdf_page(doc, i)
                units.append(PageUnit(path.name, f"page {i + 1}", ocr_text, "ocr"))
        doc.close()
    return units


def _ocr_pdf_page(doc: pymupdf.Document, page_index: int, dpi: int = 300) -> str:
    page = doc[page_index]
    pix = page.get_pixmap(dpi=dpi)
    img = Image.frombytes("RGB", (pix.width, pix.height), pix.samples)
    return pytesseract.image_to_string(img, lang="eng").strip()


def _parse_docx(path: Path) -> list[PageUnit]:
    doc = DocxDocument(path)
    text = "\n".join(p.text for p in doc.paragraphs if p.text.strip())
    # python-docx has no reliable page concept (pagination is a rendering-time
    # thing), so the whole document is one unit.
    return [PageUnit(path.name, "document", text, "native")]


def _parse_xlsx(path: Path) -> list[PageUnit]:
    wb = load_workbook(path, data_only=True)
    units = []
    for sheet in wb.worksheets:
        rows = []
        for row in sheet.iter_rows(values_only=True):
            cells = [str(c) for c in row if c is not None]
            if cells:
                rows.append(" | ".join(cells))
        if rows:
            units.append(PageUnit(path.name, f"sheet {sheet.title}", "\n".join(rows), "native"))
    return units


def _parse_pptx(path: Path) -> list[PageUnit]:
    prs = Presentation(path)
    units = []
    for i, slide in enumerate(prs.slides):
        texts = []
        for shape in slide.shapes:
            if shape.has_text_frame and shape.text_frame.text.strip():
                texts.append(shape.text_frame.text.strip())
        if texts:
            units.append(PageUnit(path.name, f"slide {i + 1}", "\n".join(texts), "native"))
    return units
